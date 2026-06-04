from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from fastapi import HTTPException
from sqlalchemy import select

from app.db.models import GradingTask, UploadedFile
from app.db.session import SessionLocal
from app.models.file import FileRole
from app.models.task import GradingTaskStatus
from app.services.question_analyzer_service import question_analyzer_service


SUPPORTED_PARSE_SUFFIXES = {".docx", ".pdf", ".txt", ".md"}


@dataclass
class ParsedFileSummary:
    id: int
    file_name: str
    file_role: str
    parsed_text_length: int


class FileParseService:
    def parse_task_files(self, task_id: int) -> dict:
        with SessionLocal() as db:
            task = db.get(GradingTask, task_id)
            if task is None:
                raise HTTPException(status_code=404, detail="Task not found")

            files = db.scalars(
                select(UploadedFile)
                .where(UploadedFile.task_id == task_id)
                .order_by(UploadedFile.file_role, UploadedFile.id),
            ).all()
            if not files:
                raise HTTPException(status_code=400, detail="No files uploaded for this task")

            self._ensure_required_materials(files)

            parsed_files: list[ParsedFileSummary] = []
            failed_files: list[dict] = []
            for uploaded_file in files:
                suffix = Path(uploaded_file.storage_path).suffix.lower()
                if suffix == ".zip":
                    continue
                try:
                    parsed_text = parse_document_text(Path(uploaded_file.storage_path))
                    uploaded_file.parsed_text = parsed_text
                    parsed_files.append(
                        ParsedFileSummary(
                            id=uploaded_file.id,
                            file_name=uploaded_file.file_name,
                            file_role=uploaded_file.file_role.value,
                            parsed_text_length=len(parsed_text),
                        ),
                    )
                except ValueError as exc:
                    failed_files.append(
                        {
                            "id": uploaded_file.id,
                            "file_name": uploaded_file.file_name,
                            "file_role": uploaded_file.file_role.value,
                            "error": str(exc),
                        },
                    )

            if not parsed_files:
                raise HTTPException(status_code=400, detail="No supported files could be parsed")

            question_analyzer_service._clear_after_question_analysis(db, task_id)
            task.status = GradingTaskStatus.PARSED
            db.commit()

            return {
                "task_id": task_id,
                "status": task.status.value,
                "stage": "parse_files",
                "parsed_count": len(parsed_files),
                "failed_count": len(failed_files),
                "parsed_files": [summary.__dict__ for summary in parsed_files],
                "failed_files": failed_files,
            }

    def _ensure_required_materials(self, files: list[UploadedFile]) -> None:
        roles = {uploaded_file.file_role for uploaded_file in files}
        missing_roles = []
        if FileRole.REQUIREMENT not in roles:
            missing_roles.append("作业要求")
        if FileRole.REFERENCE_ANSWER not in roles:
            missing_roles.append("参考答案")
        if missing_roles:
            raise HTTPException(status_code=400, detail=f"请先上传：{'、'.join(missing_roles)}")


def parse_document_text(path: Path) -> str:
    if not path.exists():
        raise ValueError("文件不存在")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_PARSE_SUFFIXES:
        raise ValueError(f"暂不支持解析 {suffix or '无扩展名'} 文件")

    try:
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            table_cells = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
                if cell.text.strip()
            ]
            return "\n".join([*paragraphs, *table_cells])
        if suffix == ".pdf":
            import fitz

            with fitz.open(path) as document:
                return "\n".join(page.get_text() for page in document)
    except Exception as exc:
        raise ValueError(f"文件解析失败：{exc}") from exc

    raise ValueError(f"暂不支持解析 {suffix or '无扩展名'} 文件")


file_parse_service = FileParseService()
