from __future__ import annotations

from io import BytesIO
from datetime import datetime, timezone
from typing import Optional

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.db.models import ClassGroup, ClassStudent
from app.db.session import SessionLocal
from app.schemas.class_group import (
    ClassGroupCreate,
    ClassGroupRead,
    ClassGroupUpdate,
    ClassStudentCreate,
    ClassStudentRead,
    ClassStudentUpdate,
    StudentImportResult,
)


NAME_HEADERS = {"姓名", "学生姓名", "名字", "名称", "studentname", "name"}
NO_HEADERS = {"学号", "学生学号", "学生编号", "编号", "studentno", "studentid", "no"}


class ClassGroupService:
    def create_class_group(self, payload: ClassGroupCreate) -> ClassGroupRead:
        with SessionLocal() as db:
            class_group = ClassGroup(
                class_name=payload.class_name,
                note=payload.note,
                student_count=0,
            )
            db.add(class_group)
            db.commit()
            db.refresh(class_group)
            return self._to_read_schema(class_group)

    def list_class_groups(self) -> list[ClassGroupRead]:
        with SessionLocal() as db:
            class_groups = db.scalars(select(ClassGroup).order_by(ClassGroup.id.desc())).all()
            return [self._to_read_schema(class_group) for class_group in class_groups]

    def get_class_group(self, class_group_id: int) -> Optional[ClassGroupRead]:
        with SessionLocal() as db:
            class_group = db.get(ClassGroup, class_group_id)
            if class_group is None:
                return None
            return self._to_read_schema(class_group)

    def update_class_group(self, class_group_id: int, payload: ClassGroupUpdate) -> ClassGroupRead:
        with SessionLocal() as db:
            class_group = db.get(ClassGroup, class_group_id)
            if class_group is None:
                raise HTTPException(status_code=404, detail="Class not found")
            update_data = payload.model_dump(exclude_unset=True)
            for key, value in update_data.items():
                setattr(class_group, key, value)
            db.commit()
            db.refresh(class_group)
            return self._to_read_schema(class_group)

    def delete_class_group(self, class_group_id: int) -> None:
        with SessionLocal() as db:
            class_group = db.get(ClassGroup, class_group_id)
            if class_group is None:
                raise HTTPException(status_code=404, detail="Class not found")
            db.delete(class_group)
            db.commit()

    def list_students(self, class_group_id: int) -> list[ClassStudentRead]:
        with SessionLocal() as db:
            self._get_class_group_or_404(db, class_group_id)
            students = db.scalars(
                select(ClassStudent)
                .where(ClassStudent.class_group_id == class_group_id)
                .order_by(ClassStudent.id.desc()),
            ).all()
            return [self._to_student_schema(student) for student in students]

    def create_student(self, class_group_id: int, payload: ClassStudentCreate) -> ClassStudentRead:
        with SessionLocal() as db:
            self._get_class_group_or_404(db, class_group_id)
            self._ensure_student_not_duplicate(db, class_group_id, payload.student_no, exclude_student_id=None)
            student = ClassStudent(
                class_group_id=class_group_id,
                student_name=payload.student_name.strip(),
                student_no=payload.student_no.strip(),
            )
            db.add(student)
            db.flush()
            self._sync_student_count(db, class_group_id)
            db.commit()
            db.refresh(student)
            return self._to_student_schema(student)

    def update_student(
        self,
        class_group_id: int,
        student_id: int,
        payload: ClassStudentUpdate,
    ) -> ClassStudentRead:
        with SessionLocal() as db:
            self._get_class_group_or_404(db, class_group_id)
            student = self._get_student_or_404(db, class_group_id, student_id)
            update_data = payload.model_dump(exclude_unset=True)
            if "student_no" in update_data:
                self._ensure_student_not_duplicate(
                    db,
                    class_group_id,
                    update_data["student_no"] or "",
                    exclude_student_id=student_id,
                )
            for key, value in update_data.items():
                setattr(student, key, value.strip() if isinstance(value, str) else value)
            db.commit()
            db.refresh(student)
            return self._to_student_schema(student)

    def delete_student(self, class_group_id: int, student_id: int) -> None:
        with SessionLocal() as db:
            self._get_class_group_or_404(db, class_group_id)
            student = self._get_student_or_404(db, class_group_id, student_id)
            db.delete(student)
            db.flush()
            self._sync_student_count(db, class_group_id)
            db.commit()

    async def import_students_from_excel(self, class_group_id: int, file: UploadFile) -> StudentImportResult:
        if not file.filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=400, detail="Only .xlsx files are supported")

        content = await file.read()
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail="Invalid Excel file") from exc

        rows, errors = self._parse_student_rows(workbook)
        if not rows and not errors:
            errors.append("No valid student rows found")

        with SessionLocal() as db:
            self._get_class_group_or_404(db, class_group_id)
            existing_students = db.scalars(
                select(ClassStudent).where(ClassStudent.class_group_id == class_group_id),
            ).all()
            existing_numbers = {item.student_no for item in existing_students if item.student_no}
            existing_names_without_no = {
                item.student_name for item in existing_students if not item.student_no
            }

            imported: list[ClassStudent] = []
            skipped_count = 0
            for row in rows:
                student_name = row["student_name"].strip()
                student_no = row["student_no"].strip()
                if student_no and student_no in existing_numbers:
                    skipped_count += 1
                    continue
                if not student_no and student_name in existing_names_without_no:
                    skipped_count += 1
                    continue
                student = ClassStudent(
                    class_group_id=class_group_id,
                    student_name=student_name,
                    student_no=student_no,
                )
                db.add(student)
                imported.append(student)
                if student_no:
                    existing_numbers.add(student_no)
                else:
                    existing_names_without_no.add(student_name)

            db.flush()
            self._sync_student_count(db, class_group_id)
            db.commit()
            for student in imported:
                db.refresh(student)

            return StudentImportResult(
                imported_count=len(imported),
                skipped_count=skipped_count,
                students=[self._to_student_schema(student) for student in imported],
                errors=errors,
            )

    def _to_read_schema(self, class_group: ClassGroup) -> ClassGroupRead:
        return ClassGroupRead(
            id=class_group.id,
            class_name=class_group.class_name,
            note=class_group.note,
            student_count=class_group.student_count,
            created_at=self._ensure_datetime(class_group.created_at),
            updated_at=self._ensure_datetime(class_group.updated_at),
        )

    def _ensure_datetime(self, value: datetime) -> datetime:
        if value.tzinfo is None:
            return value.replace(tzinfo=timezone.utc)
        return value

    def _to_student_schema(self, student: ClassStudent) -> ClassStudentRead:
        return ClassStudentRead(
            id=student.id,
            class_group_id=student.class_group_id,
            student_name=student.student_name,
            student_no=student.student_no,
            created_at=self._ensure_datetime(student.created_at),
            updated_at=self._ensure_datetime(student.updated_at),
        )

    def _get_class_group_or_404(self, db, class_group_id: int) -> ClassGroup:
        class_group = db.get(ClassGroup, class_group_id)
        if class_group is None:
            raise HTTPException(status_code=404, detail="Class not found")
        return class_group

    def _get_student_or_404(self, db, class_group_id: int, student_id: int) -> ClassStudent:
        student = db.get(ClassStudent, student_id)
        if student is None or student.class_group_id != class_group_id:
            raise HTTPException(status_code=404, detail="Student not found")
        return student

    def _ensure_student_not_duplicate(
        self,
        db,
        class_group_id: int,
        student_no: str,
        exclude_student_id: int | None,
    ) -> None:
        student_no = student_no.strip()
        if not student_no:
            return
        statement = select(ClassStudent).where(
            ClassStudent.class_group_id == class_group_id,
            ClassStudent.student_no == student_no,
        )
        matched = db.scalars(statement).first()
        if matched is not None and matched.id != exclude_student_id:
            raise HTTPException(status_code=409, detail="Student number already exists in this class")

    def _sync_student_count(self, db, class_group_id: int) -> None:
        class_group = self._get_class_group_or_404(db, class_group_id)
        class_group.student_count = db.scalar(
            select(func.count()).select_from(ClassStudent).where(
                ClassStudent.class_group_id == class_group_id,
            ),
        ) or 0

    def _parse_student_rows(self, workbook) -> tuple[list[dict[str, str]], list[str]]:
        sheet = workbook[workbook.sheetnames[0]]
        if hasattr(sheet, "reset_dimensions"):
            sheet.reset_dimensions()
        raw_rows = [
            [self._cell_to_text(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        raw_rows = [row for row in raw_rows if any(row)]
        if not raw_rows:
            return [], ["Excel file is empty"]

        header = [self._normalize_header(value) for value in raw_rows[0]]
        name_index = self._find_header_index(header, NAME_HEADERS)
        no_index = self._find_header_index(header, NO_HEADERS)
        start_index = 1 if name_index is not None else 0
        if name_index is None:
            name_index = 0
            no_index = 1 if len(raw_rows[0]) > 1 else None

        rows: list[dict[str, str]] = []
        errors: list[str] = []
        for row_number, row in enumerate(raw_rows[start_index:], start=start_index + 1):
            student_name = row[name_index].strip() if len(row) > name_index else ""
            student_no = row[no_index].strip() if no_index is not None and len(row) > no_index else ""
            if not student_name:
                if any(row):
                    errors.append(f"Row {row_number}: missing student name")
                continue
            rows.append({"student_name": student_name, "student_no": student_no})
        return rows, errors

    def _find_header_index(self, header: list[str], candidates: set[str]) -> int | None:
        for index, value in enumerate(header):
            if value in candidates:
                return index
        return None

    def _normalize_header(self, value: str) -> str:
        return value.strip().lower().replace(" ", "").replace("_", "")

    def _cell_to_text(self, value) -> str:
        if value is None:
            return ""
        return str(value).strip()


class_group_service = ClassGroupService()
