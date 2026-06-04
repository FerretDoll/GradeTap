from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.db.models import (
    AnswerEvidence,
    GradingDeduction,
    GradingResult,
    GradingTask,
    Question,
    StudentSubmission,
    TeacherRevision,
)
from app.db.session import SessionLocal
from app.models.grading import ReviewStatus
from app.models.task import GradingTaskStatus
from app.schemas.export import ExportTeacherRevision


def export_grades_to_excel(task_id: int, teacher_revisions: list[ExportTeacherRevision] | None = None) -> bytes:
    with SessionLocal() as db:
        task = db.get(GradingTask, task_id)
        if task is None:
            raise ValueError("Task not found")

        questions = db.scalars(
            select(Question)
            .where(Question.task_id == task_id)
            .order_by(Question.sort_order, Question.id),
        ).all()
        submissions = db.scalars(
            select(StudentSubmission)
            .where(StudentSubmission.task_id == task_id)
            .order_by(StudentSubmission.student_no, StudentSubmission.student_name, StudentSubmission.id),
        ).all()
        results = db.scalars(
            select(GradingResult)
            .where(GradingResult.task_id == task_id)
            .order_by(GradingResult.student_submission_id, GradingResult.question_id),
        ).all()
        if not results:
            raise ValueError("No grading results to export")

        deductions = db.scalars(
            select(GradingDeduction)
            .where(GradingDeduction.grading_result_id.in_([result.id for result in results])),
        ).all()
        evidence_items = db.scalars(
            select(AnswerEvidence)
            .where(AnswerEvidence.task_id == task_id),
        ).all()
        stored_revisions = db.scalars(
            select(TeacherRevision)
            .where(TeacherRevision.task_id == task_id)
            .order_by(TeacherRevision.grading_result_id, TeacherRevision.id.desc()),
        ).all()

        task.status = GradingTaskStatus.EXPORTED
        db.commit()

    latest_stored_revisions: dict[int, TeacherRevision] = {}
    for revision in stored_revisions:
        latest_stored_revisions.setdefault(revision.grading_result_id, revision)

    return _build_workbook_bytes(
        task=task,
        questions=questions,
        submissions=submissions,
        results=results,
        deductions=deductions,
        evidence_items=evidence_items,
        teacher_revisions=[
            _revision_model_from_db(revision)
            for revision in latest_stored_revisions.values()
        ] + (teacher_revisions or []),
    )


def _build_workbook_bytes(
    *,
    task: GradingTask,
    questions: list[Question],
    submissions: list[StudentSubmission],
    results: list[GradingResult],
    deductions: list[GradingDeduction],
    evidence_items: list[AnswerEvidence],
    teacher_revisions: list[ExportTeacherRevision],
) -> bytes:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "成绩总表"
    detail_ws = wb.create_sheet("按题明细")

    revisions_by_result_id = {
        revision.grading_result_id: revision
        for revision in teacher_revisions
        if revision.grading_result_id is not None
    }
    revisions_by_answer_id = {
        revision.student_answer_id: revision
        for revision in teacher_revisions
        if revision.student_answer_id is not None
    }
    result_by_student_question = {
        (result.student_submission_id, result.question_id): result
        for result in results
    }
    deductions_by_result_id: dict[int, list[GradingDeduction]] = defaultdict(list)
    for deduction in deductions:
        deductions_by_result_id[deduction.grading_result_id].append(deduction)

    evidence_by_answer_id: dict[int, list[AnswerEvidence]] = defaultdict(list)
    for evidence in evidence_items:
        evidence_by_answer_id[evidence.student_answer_id].append(evidence)

    summary_headers = ["学号", "姓名"]
    for question in questions:
        label = _question_label(question)
        summary_headers.extend([f"{label} AI分数", f"{label} 最终分数"])
    summary_headers.extend(["总分", "是否教师修改", "是否需要复核", "复核状态", "教师评语"])
    summary_ws.append(summary_headers)

    detail_headers = [
        "学号",
        "姓名",
        "题号",
        "题目分值",
        "AI分数",
        "最终分数",
        "评分维度得分",
        "正向证据",
        "负向证据",
        "AI评语",
        "教师评语",
        "复核状态",
        "主要扣分原因",
    ]
    detail_ws.append(detail_headers)

    for submission in submissions:
        total_score = 0.0
        teacher_modified = False
        review_required = False
        review_statuses: list[str] = []
        student_comments: list[str] = []
        row: list[Any] = [submission.student_no, submission.student_name]

        for question in questions:
            result = result_by_student_question.get((submission.id, question.id))
            if result is None:
                row.extend(["", ""])
                continue

            revision = _revision_for_result(result, revisions_by_result_id, revisions_by_answer_id)
            final_score = _effective_final_score(result, revision)
            question_full_score = float(question.total_score or 0)
            final_comment = _effective_final_comment(result, revision)
            total_score += final_score
            teacher_modified = teacher_modified or _is_teacher_modified(result, revision)
            review_required = review_required or bool(result.review_required)
            review_statuses.append(_enum_value(revision.review_status if revision else result.review_status))
            if final_score < question_full_score and final_comment:
                student_comments.append(f"{_question_label(question)}：{final_comment}")

            row.extend([result.score, final_score])
            detail_ws.append(
                [
                    submission.student_no,
                    submission.student_name,
                    question.question_number,
                    question.total_score,
                    result.score,
                    final_score,
                    _format_dimension_scores(result.dimension_scores),
                    _format_evidence(evidence_by_answer_id[result.student_answer_id], "positive_evidence"),
                    _format_evidence(evidence_by_answer_id[result.student_answer_id], "negative_evidence"),
                    result.ai_comment,
                    final_comment if final_score < question_full_score else "",
                    _enum_value(revision.review_status if revision else result.review_status),
                    _format_deductions(deductions_by_result_id[result.id]),
                ]
            )

        row.extend(
            [
                round(total_score),
                "是" if teacher_modified else "否",
                "是" if review_required else "否",
                "、".join(sorted(set(filter(None, review_statuses)))),
                "\n".join(student_comments),
            ]
        )
        summary_ws.append(row)

    _style_sheet(summary_ws)
    _style_sheet(detail_ws)
    _add_meta_sheet(wb, task, questions, len(submissions))

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _revision_for_result(
    result: GradingResult,
    revisions_by_result_id: dict[int | None, ExportTeacherRevision],
    revisions_by_answer_id: dict[int | None, ExportTeacherRevision],
) -> ExportTeacherRevision | None:
    return revisions_by_result_id.get(result.id) or revisions_by_answer_id.get(result.student_answer_id)


def _effective_final_score(result: GradingResult, revision: ExportTeacherRevision | None) -> float:
    if revision is not None:
        return float(revision.final_score)
    if result.final_score is not None:
        return float(result.final_score)
    return float(result.score)


def _is_teacher_modified(result: GradingResult, revision: ExportTeacherRevision | None) -> bool:
    if revision is None:
        return _enum_value(result.review_status) == ReviewStatus.TEACHER_MODIFIED.value
    return _enum_value(revision.review_status) == ReviewStatus.TEACHER_MODIFIED.value


def _effective_final_comment(result: GradingResult, revision: ExportTeacherRevision | None) -> str:
    if revision is not None:
        return (revision.final_comment or revision.teacher_comment or "").strip()
    return (result.final_comment or "").strip()


def _question_label(question: Question) -> str:
    raw_number = str(question.question_number or "").strip()
    number = _clean_question_number(raw_number)
    if not number and question.sort_order is not None:
        number = str(int(question.sort_order) + 1)
    return f"第{number}题" if number else f"题目{question.id}"


def _clean_question_number(raw_number: str) -> str:
    if not raw_number:
        return ""
    embedded = re.search(r"第\s*([0-9一二三四五六七八九十百]+)\s*题", raw_number)
    if embedded:
        return embedded.group(1)
    leading = re.match(r"\s*([0-9一二三四五六七八九十百]+)", raw_number)
    if leading:
        return leading.group(1)
    return raw_number.split(maxsplit=1)[0].rstrip(".、:：)")


def _enum_value(value: Any) -> str:
    return getattr(value, "value", value) or ""


def _format_dimension_scores(dimension_scores: Any) -> str:
    if not isinstance(dimension_scores, list):
        return ""
    parts = []
    for item in dimension_scores:
        if not isinstance(item, dict):
            continue
        name = item.get("dimension_name") or f"维度{item.get('rubric_id', '')}"
        score = item.get("score", "")
        max_score = item.get("max_score", "")
        reason = item.get("reason", "")
        parts.append(f"{name}: {score}/{max_score}；{reason}")
    return "\n".join(parts)


def _format_evidence(evidence_items: list[AnswerEvidence], field_name: str) -> str:
    values: list[str] = []
    for evidence in evidence_items:
        field_value = getattr(evidence, field_name) or []
        if isinstance(field_value, list):
            values.extend(str(item) for item in field_value if str(item).strip())
    return "\n".join(values)


def _format_deductions(deductions: list[GradingDeduction]) -> str:
    return "\n".join(f"{deduction.reason}（-{deduction.points}）" for deduction in deductions)


def _revision_model_from_db(revision: TeacherRevision) -> ExportTeacherRevision:
    return ExportTeacherRevision(
        grading_result_id=revision.grading_result_id,
        final_score=revision.final_score,
        final_comment=revision.final_comment,
        revision_reason=revision.revision_reason,
        review_status=revision.review_status,
    )


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="E8F0FE")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:100])
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 42)


def _add_meta_sheet(wb: Workbook, task: GradingTask, questions: list[Question], student_count: int) -> None:
    ws = wb.create_sheet("导出说明")
    rows = [
        ("任务名称", task.task_name),
        ("课程", task.course_name),
        ("班级", task.class_name),
        ("题目数量", len(questions)),
        ("学生数量", student_count),
        ("分数规则", "最终分优先使用教师复核分，没有最终分时使用 AI 分数。"),
        ("评语规则", "学生总表教师评语只整合最终分未满分题目的教师评语，格式为“第x题：教师评语”。"),
    ]
    for row in rows:
        ws.append(row)
    _style_sheet(ws)
